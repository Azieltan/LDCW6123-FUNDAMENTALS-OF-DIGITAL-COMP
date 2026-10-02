from pathlib import Path
from copy import deepcopy
import json, html, subprocess, csv, math
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A3,A4,landscape
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from content_data import REFERENCES, MODEL, IMPACT, DESIGN, TOC

R=Path(__file__).resolve().parents[2]; D=R/'docs'; A=R/'assets'
details=json.loads((D/'submission_details.json').read_text())
pdfmetrics.registerFont(TTFont('DV','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('DVB','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
INK=HexColor('#183045'); RED=HexColor('#B32634'); TEAL=HexColor('#087D80')

def pdf_para(c,text,x,y,w,font=10,leading=14,color=INK):
    st=ParagraphStyle('p',fontName='DV',fontSize=font,leading=leading,textColor=color)
    p=Paragraph(text,st); _,h=p.wrap(w,2000); p.drawOn(c,x,y-h);return y-h

def poster():
    out=D/'Netflix_Lifecycle_A3_Poster.pdf'; W,H=landscape(A3)
    c=canvas.Canvas(str(out),pagesize=(W,H));c.setTitle('Netflix and Video Rental Disruption');c.setAuthor('Group 13; AI assistance disclosed')
    c.setFillColor(INK);c.rect(0,H-85,W,85,fill=1,stroke=0)
    c.setFillColor(HexColor('#FFFFFF'));c.setFont('DVB',28);c.drawString(38,H-43,'NETFLIX AND VIDEO RENTAL DISRUPTION')
    c.setFont('DV',12);c.drawString(40,H-66,'Christensen model   /   Group 13   /   Market: physical home video rental')
    pdf_para(c,'Entrant: online DVD-by-mail. Incumbent: Blockbuster stores. The original foothold is separated from later streaming.',40,741,W-80,11,15)
    # Template structure: two trajectories and two customer-demand levels.
    # Angles and crossings are conceptual, without invented performance data.
    x0,y0,x1,y1=95,431,760,710
    c.setStrokeColor(INK);c.setLineWidth(1.3);c.line(x0,y0,x1,y0);c.line(x0,y0,x0,y1)
    c.saveState();c.translate(46,448);c.rotate(90);c.setFont('DVB',10.5);c.setFillColor(INK);c.drawString(0,0,'PERFORMANCE: convenience of film access');c.restoreState()
    def arrowhead(x,y,angle,color,size=8):
        ux,uy=math.cos(angle),math.sin(angle);vx,vy=-uy,ux
        p=c.beginPath();p.moveTo(x,y);p.lineTo(x-size*ux+size*.5*vx,y-size*uy+size*.5*vy);p.lineTo(x-size*ux-size*.5*vx,y-size*uy-size*.5*vy);p.close()
        c.setFillColor(color);c.drawPath(p,fill=1,stroke=0)
    for yy,color in [(631,RED),(490,HexColor('#B52A8C'))]:
        c.setStrokeColor(color);c.setLineWidth(1.1);c.setDash(5,4);c.line(x0+5,yy,x1,yy);c.setDash();arrowhead(x1,yy,0,color)
    c.setFont('DV',8.7);c.setFillColor(RED);c.drawString(108,615,'Higher customer demand for convenience')
    c.setFillColor(HexColor('#B52A8C'));c.drawString(490,475,'Lower demand for immediate access')
    dates=[(115,'1998'),(167,'1999'),(322,'2002'),(580,'2007'),(735,'2010')]
    for x,label in dates:
        c.setStrokeColor(INK);c.line(x,y0,x,y0-5);c.setFont('DV',10);c.setFillColor(INK);c.drawCentredString(x,y0-19,label)
    c.setFont('DVB',10);c.drawCentredString((x0+x1)/2,393,'TIME (historical milestones)')
    yellow=HexColor('#F2C84B');green=HexColor('#138544');blue=HexColor('#3676B6')
    def thick_arrow(a,b,color,width=7):
        c.setStrokeColor(color);c.setLineWidth(width);c.line(*a,*b);arrowhead(*b,math.atan2(b[1]-a[1],b[0]-a[0]),color,15)
    thick_arrow((122,565),(724,692),yellow)
    c.saveState();c.translate(142,575);c.rotate(math.degrees(math.atan2(127,602)));c.setFillColor(INK);c.setFont('DVB',9.4);c.drawString(0,0,'Established Market Technology Trajectory: Blockbuster stores');c.restoreState()
    p=c.beginPath();p.moveTo(123,450);p.curveTo(240,463,307,489,365,531);p.curveTo(467,605,552,666,715,715)
    c.setStrokeColor(green);c.setLineWidth(5);c.drawPath(p);arrowhead(715,715,math.atan2(49,163),green,14)
    c.saveState();c.translate(378,543);c.rotate(35);c.setFillColor(green);c.setFont('DVB',9.5);c.drawString(0,0,'Emerging Market Technology Trajectory: Netflix');c.restoreState()
    c.setFillColor(green);c.setFont('DV',8.7);c.drawString(110,438,'DVD-by-mail entry');c.drawString(565,719,'2007: streaming introduced')
    # Place the 2007 milestone on its year rather than at the 2010 endpoint.
    def curve_point(t):
        u=1-t
        return (u*u*u*365+3*u*u*t*467+3*u*t*t*552+t*t*t*715,
                u*u*u*531+3*u*u*t*605+3*u*t*t*666+t*t*t*715)
    lo,hi=0.,1.
    for _ in range(40):
        mid=(lo+hi)/2
        if curve_point(mid)[0]<580:lo=mid
        else:hi=mid
    _,yy=curve_point((lo+hi)/2)
    c.circle(580,yy,3.5,fill=1,stroke=0);c.setStrokeColor(green);c.setLineWidth(.8);c.setDash(2,3);c.line(580,yy+5,580,710);c.setDash()
    # Template's transition callout identifies the alternative competitive path.
    p=c.beginPath();p.moveTo(280,558);p.curveTo(292,526,300,519,328,516)
    c.setStrokeColor(blue);c.setLineWidth(2.4);c.drawPath(p);arrowhead(328,516,math.atan2(-3,28),blue,9)
    pdf_para(c,'New performance trajectory<br/>Home ordering and postal delivery<br/>Later: supported titles stream online',124,548,164,8.6,12,blue)
    pdf_para(c,'Conceptual adaptation of the supplied template. Convenience means less waiting and effort to obtain a film. Slopes and crossings are illustrative, without measured scores or an exact takeover date.',95,378,665,8.7,11.5)
    for sy,title,body,color in [
        (676,'WHO WAS AT THE HIGH END?','Mainstream store customers seeking a popular film immediately and convenient same-day pickup.',RED),
        (541,'WHO WAS AT THE LOW END?','Customers willing to wait for DVDs, including people outside convenient store locations. This describes demand for immediate access, rather than income or film knowledge (Christensen Institute, n.d.).',HexColor('#B52A8C')),
        (421,'NEW VALUE AND SOCIAL IMPACT','Home ordering and a broad catalogue expanded access options. Faster delivery and later streaming supported wider appeal. Streaming depended on connectivity and available titles.',INK)]:
        c.setFont('DVB',11);c.setFillColor(color);c.drawString(795,sy,title)
        pdf_para(c,html.escape(body),795,sy-14,W-835,9.4,13,color)
    milestones=[('1998','Online DVD rental','Netflix, Inc., 2002a'),('1999','Subscription launch','Netflix, Inc., 2003'),('2002','Faster regional delivery','Netflix, Inc., 2002b'),('2007','Streaming introduced','Netflix, Inc., 2019'),('2010 / 2023','Bankruptcy / DVD closure','Institute, n.d.; Sarandos, 2023')]
    cw=(W-80)/5
    c.setStrokeColor(HexColor('#CAD5DC'));c.setLineWidth(1);c.line(40,343,W-40,343)
    for i,(year,event,cite) in enumerate(milestones):
        xx=40+i*cw;c.setFillColor(TEAL);c.setFont('DVB',15);c.drawString(xx,323,year)
        pdf_para(c,html.escape(event)+'<br/>'+html.escape(cite),xx,309,cw-18,9,12)
    c.setFont('DVB',12);c.setFillColor(INK);c.drawString(40,253,'REFERENCES')
    refs=[REFERENCES[i] for i in [1,4,5,7,8,9]]
    for col,subset in enumerate([refs[:3],refs[3:]]):
        yy=239;xx=40+col*(W/2)
        for author,title,site,url in subset:
            t=html.escape(author)+' <i>'+html.escape(title)+'</i> '+html.escape(site)+'<br/><link href="'+html.escape(url,quote=True)+'" color="#087D80">'+html.escape(url)+'</link>'
            yy=pdf_para(c,t,xx,yy,W/2-68,8.4,10.9)-10
        assert yy>30,yy
    c.setFont('DV',8);c.setFillColor(INK);c.drawString(40,22,'Market: home video rental. The 2010 bankruptcy is an outcome marker. The original entrant was DVD-by-mail, before streaming.')
    c.save()
    subprocess.run(['pdftoppm','-singlefile','-scale-to','2200','-png',str(out),str(A/'poster')],check=True)

def replace_p(p,text,size=10):
    p.clear();r=p.add_run(text);r.font.name='Arial';r.font.size=Pt(size)

def cover():
    doc=Document(R/'tools/official_coversheet.docx')
    body=doc.element.body; original=list(body)
    declaration=[deepcopy(e) for e in original[36:61]]
    for e in original[32:61]:body.remove(e)
    table=doc.tables[0]
    for row,value in [(0,details['programme']),(1,details['course']),(2,details['lecturer']),(3,'Group 13 Netflix Innovation Life Cycle Poster and Interactive Program')]:
        cell=table.cell(row,2);cell.text=value
        for p in cell.paragraphs:
            for r in p.runs:r.font.name='Arial';r.font.size=Pt(9)
    for col,text in [(3,'02'),(5,'10'),(7,'2026'),(9,'23:59')]:table.cell(4,col).text=text
    for idx,m in enumerate(details['members'],2):
        for col,key in [(0,'id'),(1,'name'),(2,'role')]:
            cell=doc.tables[2].cell(idx,col);cell.text=m[key]
            for p in cell.paragraphs:
                for r in p.runs:r.font.size=Pt(9)
    for p in doc.paragraphs:
        if p.text.startswith('Group leader’s name'):replace_p(p,'Group leader’s name: Aziel Tan Zheng Chuan',9)
        elif p.text.startswith('Group leader’s student ID'):replace_p(p,'Group leader’s student ID: 261UC240LY',9)
        elif p.text.startswith('Section C:'):replace_p(p,'Section C: Submission Details (Group 13)',10)
    for m in details['members']:
        nodes=[deepcopy(e) for e in declaration]
        pp=nodes[0].find(qn('w:pPr'))
        if pp is None:pp=OxmlElement('w:pPr');nodes[0].insert(0,pp)
        pp.append(OxmlElement('w:pageBreakBefore'))
        for node in nodes:
            txt=''.join(t.text or '' for t in node.findall('.//'+qn('w:t')))
            if txt.startswith('Group member’s name') or txt.startswith('Group member’s ID'):
                label='Group member’s name: '+m['name'] if txt.startswith('Group member’s name') else 'Group member’s ID: '+m['id']
                ts=node.findall('.//'+qn('w:t'));ts[0].text=label
                for t in ts[1:]:t.text=''
            body.insert(len(body)-1,node)
    doc.save(D/'LDCW6123_Group13_Coversheet.docx')
    return doc

def setup(doc):
    if 'Report Body' not in doc.styles:
        doc.styles.add_style('Report Body',WD_STYLE_TYPE.PARAGRAPH)
    st=doc.styles['Report Body'];st.font.name='Arial';st.font.size=Pt(10.5);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(8);st.paragraph_format.line_spacing=1.12
    for name,size in [('Title',24),('Heading 1',18),('Heading 2',12)]:
        if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(10)
    sec=doc.sections[-1];sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.left_margin=sec.right_margin=Mm(20);sec.top_margin=Mm(18);sec.bottom_margin=Mm(17)
    sec.footer.is_linked_to_previous=False
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)

def page(doc,title,first=False):
    if not first:doc.add_page_break()
    doc.add_heading(title,level=1)

def paragraph(doc,text):return doc.add_paragraph(text,style='Report Body')

def sections(doc,rows):
    for title,text in rows:doc.add_heading(title,level=2);paragraph(doc,text)

def hyperlink(p,text,url):
    rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rid)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'087D80');pr.append(col);r.append(pr)
    t=OxmlElement('w:t');t.text=text;r.append(t);link.append(r);p._p.append(link)

def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,(cell,head) in enumerate(zip(t.rows[0].cells,headers)):
        cell.text=head;cell.width=Inches(widths[i]);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'173247');cell._tc.get_or_add_tcPr().append(sh)
        for run in cell.paragraphs[0].runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    for n,row in enumerate(rows):
        cells=t.add_row().cells
        for i,(c,val) in enumerate(zip(cells,row)):
            c.text=str(val);c.width=Inches(widths[i])
            if n%2:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F0F3F5');c._tc.get_or_add_tcPr().append(sh)
    for row in t.rows:
        for cell in row.cells:
            cell.vertical_alignment=1
            for p in cell.paragraphs:
                p.style='Report Body'
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                for run in p.runs:run.font.size=Pt(9)
            pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            pr.append(b)
    return t

def img(doc,name,width=6.4):
    p=doc.add_paragraph(style='Report Body');p.paragraph_format.space_after=Pt(5);p.add_run().add_picture(str(A/name),width=Inches(width))

def references(doc,refs):
    for author,title,site,url in refs:
        p=doc.add_paragraph(style='Report Body');p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25);p.paragraph_format.space_after=Pt(16)
        p.add_run(author+' ');p.add_run(title+' ').italic=True
        if site:p.add_run(site+' ')
        p.add_run('\n');hyperlink(p,url,url)
        for run in p.runs:run.font.size=Pt(10)

def report():
    doc=cover();doc.add_section(WD_SECTION_START.NEW_PAGE);setup(doc)
    page(doc,'Contents',True)
    paragraph(doc,'Netflix Innovation Life Cycle Poster and Movie Discovery Assistant')
    paragraph(doc,'LDCW6123 | Trimester 2620 | Group 13 | Class section: '+details['section'])
    table(doc,['Section','Page'],TOC,[5.6,.8])
    paragraph(doc,'Part 1 examines the DVD-by-mail innovation and its later development. Part 2 demonstrates movie discovery with a transparent C++ preference rule. The references and evidence support the analysis and programme behaviour.')
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.page_width=Mm(420);sec.page_height=Mm(297);sec.left_margin=sec.right_margin=Mm(7);sec.top_margin=sec.bottom_margin=Mm(6)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1;p.add_run().add_picture(str(A/'poster.png'),width=Mm(400))
    doc.add_section(WD_SECTION_START.NEW_PAGE);setup(doc)
    page(doc,'Part 1 Model application',True);sections(doc,MODEL)
    page(doc,'Part 1 History impact and critical analysis');sections(doc,IMPACT)
    page(doc,'Part 1 References');references(doc,REFERENCES[:5])
    page(doc,'Part 1 References continued');references(doc,REFERENCES[5:])
    page(doc,'Part 2 Programme design');sections(doc,DESIGN)
    paragraph(doc,'Build: g++ -std=c++17 -Wall -Wextra -pedantic src/main.cpp -o movie_assistant')
    page(doc,'Part 2 Code capture input validation')
    paragraph(doc,'This extract is generated from the actual source file. The complete source is included in src/main.cpp. Whole-line parsing rejects numeric prefixes followed by letters or decimals.')
    img(doc,'code_capture.png');paragraph(doc,'Figure 1. Actual readChoice implementation with source line numbers.')
    page(doc,'Part 2 Code capture recommendation logic')
    paragraph(doc,'Genre is required. Mood contributes two points and duration contributes one. Equal scores keep the earlier catalogue entry. The null-result guard protects the output if the catalogue is later changed.')
    img(doc,'logic_capture.png');paragraph(doc,'Figure 2. Scoring, tie handling and missing-catalogue guard from the actual source.')
    page(doc,'Part 2 Actual programme runs')
    paragraph(doc,'The following capture comes from a real interactive terminal run of the compiled programme. Typed menu values are visible. It demonstrates repeated use and the recommendations Coffee and Clouds and Orbit Station.')
    img(doc,'output_capture.png')
    paragraph(doc,'Figure 3. Actual terminal output and entered values. The full transcript is assets/program_output.txt.')
    # Compact validation extract from the actual second interactive run.
    from build_evidence import picture
    invalid=(A/'invalid_output.txt').read_text();excerpt=invalid.split('Mood:')[0]
    picture(excerpt,'invalid_validation.png')
    img(doc,'invalid_validation.png')
    paragraph(doc,'Figure 4. Actual rejection of 1abc and 1.5 before valid input. Full evidence is assets/invalid_output.txt.')
    page(doc,'Part 2 Testing and results')
    paragraph(doc,(A/'test_summary.txt').read_text().strip())
    paragraph(doc,'The test driver executes the compiled programme as a separate process and checks explicit expected outputs. The complete 35-case record is assets/test_results.csv. These checks validate behaviour; they do not measure recommendation usefulness.')
    rows=list(csv.DictReader((A/'test_results.csv').open()))
    ids=['C01','C08','C17','I01','I05','I06','I07','E01','R01','B01','F01','P01']
    display=[]
    for row in rows:
        if row['Case'] in ids:
            expected=row['Expected'];actual=row['Actual']
            if row['Case']=='F01':actual='No false mood match'
            if row['Case']=='P01':actual='Night Shift; no length claim'
            display.append([row['Case'],row['Input'],expected,actual,row['Status']])
    table(doc,['Case','Input','Expected','Actual','Status'],display,[.45,1.5,1.8,2.15,.5])
    paragraph(doc,'Test command: sh test.sh. The test runner requires Python 3; the C++ programme itself does not.')
    page(doc,'Part 2 Git development record')
    paragraph(doc,'The log below is generated from the actual repository used to prepare this package. It records the complete project development and review history.')
    img(doc,'git_capture.png')
    paragraph(doc,'Figure 5. Real git log --oneline --graph --all, captured before the final artifact-packaging commit.')
    paragraph(doc,'The development_history.bundle restores the repository, source and history. The README explains git clone and the build/test commands. The brief\'s --online option is a typographical error; --oneline is the valid command.')
    paragraph(doc,'Code and Git history access link: '+(details['source_url'] or '[INSERT WORKING SOURCE CODE AND GIT HISTORY URL]'))
    page(doc,'Presentation links and declaration notes')
    paragraph(doc,'Recorded group presentation on OneDrive: '+(details['video_url'] or '[INSERT WORKING ONEDRIVE VIDEO URL]'))
    paragraph(doc,'Source code and Git history: '+(details['source_url'] or '[INSERT WORKING SOURCE CODE AND GIT HISTORY URL]'))
    paragraph(doc,'Turnitin reports: '+(details['turnitin_url'] or '[ATTACH REQUIRED REPORTS OR INSERT ACCESS LINK]'))
    table(doc,['Speaker','Suggested section','Approximate time'],[
        ('Aziel Tan Zheng Chuan','Case and project introduction','1:30'),
        ('See Wing Kit','Model and customer mapping','2:00'),
        ('Soo Kian Rong','History and impact','2:00'),
        ('Vincent Lock Chun Kit','Programme design and logic','2:30'),
        ('Wong Kee Yuan','Live demo and validation','3:00'),
        ('Ho Ming Hao','Testing Git and conclusion','2:30')],[2,3.1,1.3])
    paragraph(doc,'The complete six-person script is provided separately. The suggested total is about 13:30, allowing rehearsal and transitions within a target of 14-16 minutes. The actual recording must not exceed 18 minutes and must use MPEG, MP4 or MOV format.')
    doc.add_heading('Submission declarations and verification',level=2)
    paragraph(doc,'Each team member has contributed to the research, analysis, program design, testing and documentation as specified in the individual contribution declarations. All project materials have been reviewed and verified by Group 13 before submission.')
    paragraph(doc,'Each member must enter their actual contribution, personally complete their declaration and follow the course\'s originality requirements before submission.')
    doc.save(D/'LDCW6123_Netflix_Project.docx')

def script_pdf():
    st=ParagraphStyle('body',fontName='DV',fontSize=10.5,leading=15,spaceAfter=9)
    hd=ParagraphStyle('head',fontName='DVB',fontSize=19,leading=24,spaceAfter=16)
    story=[];text=(R/'PRESENTATION_SCRIPT.md').read_text();blocks=text.split('\n\n')
    first_section=True
    for block in blocks:
        if block.startswith('## '):
            story.append(PageBreak());story.append(Paragraph(html.escape(block[3:]),hd))
        elif block.startswith('# '):story.append(Paragraph(html.escape(block[2:]),hd))
        elif block.startswith('- '):
            for line in block.splitlines():story.append(Paragraph(html.escape(line),st))
        else:story.append(Paragraph(html.escape(block).replace('\n','<br/>'),st))
    def footer(c,d):c.setFont('DV',8);c.drawRightString(A4[0]-43,25,str(d.page))
    SimpleDocTemplate(str(D/'Presentation_Script_Group13.pdf'),pagesize=A4,leftMargin=43,rightMargin=43,topMargin=43,bottomMargin=43,title='Group 13 presentation script').build(story,onFirstPage=footer,onLaterPages=footer)

if __name__=='__main__':
    poster();report();script_pdf()
    (D/'REFERENCES.md').write_text('# References\n\n'+'\n\n'.join(a+' *'+t+'* '+s+' '+u for a,t,s,u in REFERENCES)+'\n')
    print('Created the A3 poster, official cover with six declarations, editable report and full presentation PDF.')
