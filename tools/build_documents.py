"""Build the Spotify A3 poster, group report, reference list and script PDF."""
from pathlib import Path
import csv, html, json, subprocess
from importlib.util import module_from_spec, spec_from_file_location
from docx.enum.section import WD_SECTION_START
from docx.shared import Mm
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, PageBreak, SimpleDocTemplate
from reportlab.pdfgen import canvas

from content_data import REFERENCES, MODEL, IMPACT, DESIGN, TOC
from spotify_graph import draw_graph

R = Path(__file__).resolve().parents[1]
D = R/'docs'
A = R/'assets'
spec = spec_from_file_location('legacy_builder', R/'legacy_netflix/tools/build_documents.py')
base = module_from_spec(spec)
spec.loader.exec_module(base)
# Reuse the original MMU coversheet and Word layout helpers in this project.
base.R, base.D, base.A = R, D, A
DETAILS = json.loads((D/'submission_details.json').read_text())
INK = HexColor('#183045')
GREEN = HexColor('#087D67')
PINK = HexColor('#A43191')
RED = HexColor('#B32634')

def para(c, value, x, y, w, size=10, leading=14, color=INK):
    return base.pdf_para(c, value, x, y, w, size, leading, color)

def poster():
    draw_graph()
    dest = D/'Spotify_Clayton_A3_Poster.pdf'
    W,H = landscape(A3)
    c = canvas.Canvas(str(dest), pagesize=(W,H))
    c.setTitle('Spotify and the Disruption of Physical Music')
    c.setAuthor('LDCW6123 Group 13; AI assistance disclosed')
    c.setFillColor(INK);c.rect(0,H-85,W,85,fill=1,stroke=0)
    c.setFillColor(HexColor('#FFFFFF'));c.setFont('DVB',27)
    c.drawString(38,H-42,'FROM OWNERSHIP TO ACCESS')
    c.setFont('DV',11.5)
    c.drawString(40,H-65,'Spotify vs CDs  /  Clayton Christensen disruptive innovation model  /  LDCW6123 Group 13')
    c.drawImage(str(A/'spotify_clayton_graph.png'),38,253,width=810,height=479,mask='auto')
    # Plain prose alongside the graph makes the conventional and alternative values distinct.
    cards=[
        (721,'ESTABLISHED MARKET','A purchased CD provides a lossless 16-bit/44.1 kHz digital signal and lasting possession of that disc (Sony, n.d.).',INK),
        (627,'EMERGING ENTRANT','Spotify launched in 2008 with access to licensed music. Early compressed streaming traded conventional signal fidelity for immediacy (Spotify, n.d.).',GREEN),
        (521,'DIFFERENT VALUE','Search and discover music on demand without buying each CD. Mobile access arrived in 2009; Discover Weekly followed in 2015 (Spotify, n.d.).',INK),
        (415,'CUSTOMERS','High-end: lossless-focused CD listeners. Lower conventional demand: people willing to accept compressed playback for access. These are needs, not income groups.',PINK),
        (304,'MODEL LIMIT','The near-level CD path reflects a fixed format. The 2017 industry revenue lead is not an audio-quality crossing or proof that Spotify alone caused CD decline.',RED),
    ]
    for y,title,body,color in cards:
        c.setFillColor(color);c.setFont('DVB',10.5);c.drawString(871,y,title)
        para(c,html.escape(body),871,y-10,278,9.4,12.6)
    c.setStrokeColor(HexColor('#C8D4D8'));c.line(40,237,W-40,237)
    milestones=[
        ('2008','Spotify launches in six European markets'),
        ('2009','Mobile access'),
        ('2015','Discover Weekly'),
        ('2017','Streaming leads global recorded-music revenue*'),
        ('2025','Premium Lossless rollout; physical formats persist'),
    ]
    cw=(W-80)/5
    for n,(year,event) in enumerate(milestones):
        x=40+n*cw;c.setFillColor(GREEN);c.setFont('DVB',15);c.drawString(x,219,year)
        para(c,html.escape(event),x,208,cw-17,8.8,11.5)
    c.setFillColor(INK);c.setFont('DVB',11);c.drawString(40,151,'SELECTED REFERENCES')
    refs=[REFERENCES[i] for i in [0,1,2,4,6,7]]
    for col,subset in enumerate([refs[:3],refs[3:]]):
        x=40+col*(W/2);y=139
        for author,title,site,url in subset:
            val=html.escape(author)+' <i>'+html.escape(title)+'</i> '+html.escape(site)
            val+='<br/><link href="'+html.escape(url,quote=True)+'" color="#087D67">'+html.escape(url)+'</link>'
            y=para(c,val,x,y,W/2-65,7.55,9.2)-5
        assert y>25,(col,y)
    c.setFillColor(INK);c.setFont('DV',7.5)
    c.drawString(40,19,'*IFPI (2018) reports 38.4% of global industry revenue in 2017. Conceptual graph: vertical positions are not numerical measurements.')
    c.save()
    subprocess.run(['pdftoppm','-singlefile','-scale-to','3400','-png',str(dest),str(A/'poster')],check=True)
    return dest

def source_references(doc, subset):
    for author,title,site,url in subset:
        p=doc.add_paragraph(style='Report Body')
        p.paragraph_format.left_indent=base.Inches(.25)
        p.paragraph_format.first_line_indent=base.Inches(-.25)
        p.paragraph_format.space_after=base.Pt(15)
        p.add_run(author+' ');p.add_run(title+' ').italic=True
        if site:p.add_run(site+' ')
        if url:
            p.add_run('\n');base.hyperlink(p,url,url)

def report():
    doc=base.cover()
    # Keep the original assessment form and declaration pages; correct only its title.
    doc.tables[0].cell(3,2).text='Group 13 Spotify and Physical Music Clayton Model and Interactive Program'
    doc.save(D/'LDCW6123_Group13_Coversheet.docx')
    doc.add_section(WD_SECTION_START.NEW_PAGE);base.setup(doc)
    base.page(doc,'Contents',True)
    base.paragraph(doc,'Spotify and Physical Music Clayton Model and Music Discovery Assistant')
    base.paragraph(doc,'LDCW6123 | Trimester 2620 | Group 13 | Class section: '+DETAILS['section'])
    base.table(doc,['Section','Page'],TOC,[5.6,.8])
    base.paragraph(doc,'Part 1 applies the lecturer’s Clayton graph to CD listening and Spotify streaming with one stated audio-fidelity metric. Part 2 shows an offline, rule-based music-discovery demonstration. Historical events and a conceptual model interpretation are kept distinct.')
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.page_width=Mm(420);sec.page_height=Mm(297)
    sec.left_margin=sec.right_margin=Mm(7);sec.top_margin=sec.bottom_margin=Mm(6)
    p=doc.add_paragraph();p.paragraph_format.space_after=base.Pt(0);p.paragraph_format.line_spacing=1
    p.add_run().add_picture(str(A/'poster.png'),width=Mm(400))
    doc.add_section(WD_SECTION_START.NEW_PAGE);base.setup(doc)
    base.page(doc,'Part 1 Model application',True);base.sections(doc,MODEL)
    base.page(doc,'Part 1 Historical development and impact');base.sections(doc,IMPACT[:3])
    base.page(doc,'Part 1 Tradeoffs and critical interpretation');base.sections(doc,IMPACT[3:])
    base.page(doc,'Part 1 References');source_references(doc,REFERENCES[:4])
    base.page(doc,'Part 1 References continued');source_references(doc,REFERENCES[4:])
    base.page(doc,'Part 2 Programme design');base.sections(doc,DESIGN)
    base.paragraph(doc,'Build: g++ -std=c++17 -Wall -Wextra -pedantic src/main.cpp -o music_assistant')
    base.page(doc,'Part 2 Input validation source')
    base.paragraph(doc,'The image is generated from the actual C++ file. readChoice validates each whole line before accepting a bounded integer. Full source is included in src/main.cpp.')
    base.img(doc,'code_capture.png');base.paragraph(doc,'Figure 1. Actual readChoice source with line numbers.')
    base.page(doc,'Part 2 Recommendation source')
    base.paragraph(doc,'A required genre filter, mood weight of two and length weight of one determine each result. Equal scores retain the first track; a missing-genre guard prevents a null dereference.')
    base.img(doc,'logic_capture.png');base.paragraph(doc,'Figure 2. Actual selection code and fallback.')
    base.page(doc,'Part 2 Programme output and validation')
    base.paragraph(doc,'A genuine terminal session recommends Slow Orbit, then City Lanterns after a second search. The explanation does not claim a Reflective mood match for City Lanterns.')
    base.img(doc,'output_capture.png',6.2)
    base.paragraph(doc,'Figure 3. Interactive output and typed choices; complete transcript in assets/program_output.txt.')
    base.img(doc,'invalid_validation.png',5.2)
    base.paragraph(doc,'Figure 4. Real rejection of 1abc and 1.5; complete transcript in assets/invalid_output.txt.')
    base.page(doc,'Part 2 Testing and results')
    base.paragraph(doc,(A/'test_summary.txt').read_text().strip())
    base.paragraph(doc,'The Python driver executes the compiled programme as a separate process and writes all 35 individual outcomes to assets/test_results.csv. These checks validate deterministic behaviour rather than listener satisfaction.')
    rows=list(csv.DictReader((A/'test_results.csv').open()))
    keep={'C01','C08','C17','I01','I05','I06','I07','E01','R01','B01','F01','P01'}
    display=[]
    for row in rows:
        if row['Case'] in keep:
            actual=row['Actual']
            if row['Case']=='F01':actual='No false mood match'
            if row['Case']=='P01':actual='Storm Window; no false length claim'
            display.append([row['Case'],row['Input'],row['Expected'],actual,row['Status']])
    base.table(doc,['Case','Input','Expected','Actual','Status'],display,[.45,1.5,1.8,2.15,.5])
    base.paragraph(doc,'Run sh test.sh. Python is needed to reproduce tests, not to run the C++ programme.')
    base.page(doc,'Part 2 Git development record')
    base.paragraph(doc,'This screenshot is generated from the actual Git repository. Earlier Netflix commits remain in its history; subsequent Spotify work is honestly recorded as a later AI-assisted revision, not attributed to students who did not make those commits.')
    base.img(doc,'git_capture.png')
    base.paragraph(doc,'Figure 5. Real git log --oneline --graph --all. Names and timestamps are in assets/git_history_authors.txt. The portable repository bundle is development_history.bundle.')
    base.paragraph(doc,'Each member should review the work and add their real contribution record. The assignment’s --online example is a typo; the Git option used here is --oneline.')
    base.paragraph(doc,'Source code and Git history access link: '+(DETAILS['source_url'] or '[INSERT WORKING SOURCE CODE AND GIT HISTORY URL]'))
    base.page(doc,'Presentation links and assistance disclosure')
    base.paragraph(doc,'OneDrive video recording: '+(DETAILS['video_url'] or '[INSERT WORKING ONEDRIVE VIDEO URL]'))
    base.paragraph(doc,'Source code and Git history: '+(DETAILS['source_url'] or '[INSERT WORKING SOURCE CODE AND GIT HISTORY URL]'))
    base.paragraph(doc,'Turnitin reports: '+(DETAILS['turnitin_url'] or '[ATTACH REQUIRED REPORTS OR INSERT ACCESS LINK]'))
    base.table(doc,['Speaker','Suggested subject','Approximate time'],[
        ('Aziel Tan Zheng Chuan','Case and thesis','1:40'),('See Wing Kit','Graph and customer levels','2:20'),
        ('Soo Kian Rong','Timeline and market impact','2:15'),('Vincent Lock Chun Kit','Programme design','2:25'),
        ('Wong Kee Yuan','Demo and validation','3:00'),('Ho Ming Hao','Testing, critique and close','2:30')],[2,3.1,1.3])
    base.paragraph(doc,'The six-person script is supplied separately. The suggested total is 14:10, allowing transitions within the 18-minute maximum. The group must make its own MPEG, MP4 or MOV recording.')
    doc.add_heading('Generative AI assistance disclosure',level=2)
    base.paragraph(doc,'OpenAI ChatGPT (Codex) assisted with the prior Netflix draft on 25 September and 30 September 2026, and the Spotify revision on 1 October 2026: research, graph and document drafting, code, automated tests, evidence captures and script. The Git record shows assisted work. No member’s independent contribution, signature, completed video, Turnitin result or submission link is asserted. See AI_DISCLOSURE.txt.')
    base.paragraph(doc,'Each member must state their actual work and complete their own declaration. Confirm the class section, links, video, Turnitin reports and course originality requirements before submission.')
    dest=D/'LDCW6123_Spotify_Project.docx';doc.save(dest);return dest

def script_pdf():
    style=ParagraphStyle('body',fontName='DV',fontSize=10.5,leading=15,spaceAfter=9)
    head=ParagraphStyle('head',fontName='DVB',fontSize=18,leading=23,spaceAfter=13)
    story=[];first_speaker=True
    for block in (R/'PRESENTATION_SCRIPT.md').read_text().split('\n\n'):
        if block.startswith('## '):
            if not first_speaker:story.append(PageBreak())
            story.append(Paragraph(html.escape(block[3:]),head));first_speaker=False
        elif block.startswith('# '):story.append(Paragraph(html.escape(block[2:]),head))
        else:story.append(Paragraph(html.escape(block).replace('\n','<br/>'),style))
    def footer(c,d):c.setFont('DV',8);c.drawRightString(A4[0]-43,25,str(d.page))
    dest=D/'Presentation_Script_Group13.pdf'
    SimpleDocTemplate(str(dest),pagesize=A4,leftMargin=43,rightMargin=43,topMargin=43,bottomMargin=43,title='Group 13 Spotify presentation script').build(story,onFirstPage=footer,onLaterPages=footer)
    return dest

if __name__=='__main__':
    poster();report();script_pdf()
    (D/'REFERENCES.md').write_text('# References\n\n'+'\n\n'.join(a+' *'+t+'* '+s+(' '+u if u else '') for a,t,s,u in REFERENCES)+'\n')
    print('Created Spotify A3 poster, six-declaration coversheet, report DOCX, script PDF and references.')
